"""Application report renderer: shared layout for editable PPTX and vector PDF.
Original images are embedded with contain-fit, never cropped or enhanced.
"""
from pathlib import Path
from collections import defaultdict
import csv, json, math, hashlib, zipfile
from io import BytesIO
from datetime import datetime, timezone
from PIL import Image
from .research import *
from .classification import feature_element_ratios, criterion_rule_text

NAVY='#20334D'; MUTED='#687A91'; BLUE='#526BB5'; RED='#C15B64'; GREEN='#398E83'; AMBER='#D6A64E'; LIGHT='#F1F4F8'
def textnum(v):return '-' if v is None else f'{v:.1f}'
def percent(v):return textnum(v)+'%' if v is not None else '-'
def fit(iw,ih,x,y,w,h):
    scale=min(w/iw,h/ih);nw,nh=iw*scale,ih*scale
    return x+(w-nw)/2,y+(h-nh)/2,nw,nh

class Document:
    W,H=13.333,7.5
    def __init__(self,fmt,out,footer_text=None):
        self.fmt,self.out=fmt,Path(out);self.number=0;self.page_offset=0;self.footer_text=footer_text or "SEM / EDS screening"
        if fmt=='pptx':
            from pptx import Presentation
            from pptx.util import Inches
            self.doc=Presentation();self.doc.slide_width=Inches(self.W);self.doc.slide_height=Inches(self.H)
        else:
            from reportlab.pdfgen import canvas
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.cidfonts import UnicodeCIDFont
            pdfmetrics.registerFont(UnicodeCIDFont('HYSMyeongJo-Medium'))
            self.doc=canvas.Canvas(str(out),pagesize=(self.W*72,self.H*72),pageCompression=1)
            self.doc.setTitle('UV Tape Residue Research')
    def page(self,title,subtitle=''):
        if self.number and self.fmt=='pdf':self.doc.showPage()
        self.number+=1
        if self.fmt=='pptx':self.slide=self.doc.slides.add_slide(self.doc.slide_layouts[6])
        self.text(.35,.22,12.5,.4,title,22,bold=True)
        self.text(.35,.72,12.5,.35,subtitle,9,color=MUTED)
        self.rect(.35,1.1,12.6,.015,BLUE)
        self.text(.35,7.15,11.9,.16,self.footer_text,7,color=MUTED)
        self.text(12.35,7.13,.6,.18,str(self.number+self.page_offset),8,color=MUTED)
    def rect(self,x,y,w,h,color=LIGHT):
        if self.fmt=='pptx':
            from pptx.util import Inches
            from pptx.dml.color import RGBColor
            s=self.slide.shapes.add_shape(1,Inches(x),Inches(y),Inches(w),Inches(h));s.fill.solid();s.fill.fore_color.rgb=RGBColor.from_string(color.lstrip('#'));s.line.fill.background()
        else:
            from reportlab.lib.colors import HexColor
            self.doc.setFillColor(HexColor(color));self.doc.rect(x*72,(self.H-y-h)*72,w*72,h*72,stroke=0,fill=1)
    def text(self,x,y,w,h,value,size=12,color=NAVY,bold=False):
        value=str(value)
        if self.fmt=='pptx':
            from pptx.util import Inches,Pt
            from pptx.dml.color import RGBColor
            s=self.slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=s.text_frame;tf.word_wrap=True;tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
            for i,line in enumerate(value.split('\n')):
                p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Arial';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color.lstrip('#'))
        else:
            from reportlab.lib.colors import HexColor
            from reportlab.pdfbase.pdfmetrics import stringWidth
            font='HYSMyeongJo-Medium' if any(ord(c)>255 for c in value) else 'Helvetica-Bold' if bold else 'Helvetica'
            self.doc.setFillColor(HexColor(color));self.doc.setFont(font,size)
            yy=(self.H-y)*72-size
            for paragraph in value.split('\n'):
                line=''
                for char in paragraph:
                    if stringWidth(line+char,font,size)>w*72 and line:
                        self.doc.drawString(x*72,yy,line);yy-=size*1.25;line=''
                    line+=char
                self.doc.drawString(x*72,yy,line);yy-=size*1.25
    def image(self,path,x,y,w,h,polygons=None):
        if not path:
            self.rect(x,y,w,h,LIGHT);self.text(x+.1,y+h/2-.15,w-.2,.4,'Image unavailable',11,color=MUTED);return
        with Image.open(path) as im:iw,ih=im.size
        xx,yy,ww,hh=fit(iw,ih,x,y,w,h)
        if self.fmt=='pptx':
            from pptx.util import Inches
            self.slide.shapes.add_picture(str(path),Inches(xx),Inches(yy),width=Inches(ww),height=Inches(hh))
        else:self.doc.drawImage(str(path),xx*72,(self.H-yy-hh)*72,ww*72,hh*72,preserveAspectRatio=True,mask='auto')
        if polygons:
            if isinstance(polygons[0],dict):polygons=[polygons]
            for poly in polygons:
                pts=[]
                for p in poly if isinstance(poly,list) else []:
                    try:
                        px,py=(float(p['x']),float(p['y'])) if isinstance(p,dict) else (float(p[0]),float(p[1]))
                        if math.isfinite(px) and math.isfinite(py):pts.append((xx+max(0,min(1,px))*ww,yy+max(0,min(.92,py))*hh))
                    except (TypeError,ValueError,KeyError,IndexError):pass
                if len(pts)<3:continue
                if self.fmt=='pptx':
                    from pptx.util import Inches,Pt
                    from pptx.dml.color import RGBColor
                    from pptx.enum.shapes import MSO_CONNECTOR
                    for a,b in zip(pts,pts[1:]+pts[:1]):
                        s=self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(a[0]),Inches(a[1]),Inches(b[0]),Inches(b[1]));s.line.color.rgb=RGBColor(220,45,55);s.line.width=Pt(.8)
                else:
                    self.doc.setStrokeColorRGB(.86,.18,.22);self.doc.setLineWidth(.8);p=self.doc.beginPath();p.moveTo(pts[0][0]*72,(self.H-pts[0][1])*72)
                    for a,b in pts[1:]:p.lineTo(a*72,(self.H-b)*72)
                    p.close();self.doc.drawPath(p)
    def table(self,headers,rows,x,y,widths,row_h=.56,size=11):
        for row_i,row in enumerate([headers]+rows):
            xx=x
            for j,value in enumerate(row):
                self.rect(xx,y+row_i*row_h,widths[j]-.015,row_h-.015,LIGHT if row_i==0 else '#FAFBFD' if row_i%2 else '#FFFFFF')
                self.text(xx+.09,y+row_i*row_h+.1,widths[j]-.18,row_h-.12,value,size if row_i else size-1,bold=row_i==0);xx+=widths[j]
    def save(self):
        if self.fmt=='pptx':self.doc.save(self.out)
        else:self.doc.save()
        return self.out

class Assets:
    """Load one point at a time. No fallback from raw SEM to an overlay."""
    def __init__(self,folder,storage=None):self.folder=Path(folder);self.folder.mkdir(parents=True,exist_ok=True);self.storage=storage;self.missing=[];self.hashes={}
    def get(self,p,key):
        rid=str(p.get('id'));asset=(p.get('assets') or {}).get(key);path=None
        if isinstance(asset,str) and Path(asset).is_file():path=Path(asset)
        r2key=(p.get('r2_assets') or {}).get(key)
        if not r2key and isinstance(asset,str) and not asset.startswith(('/','http')):r2key=asset
        if path is None and r2key and self.storage is not None and self.storage.configured:
            path=self.folder/(hashlib.sha256((rid+key).encode()).hexdigest()+'.jpg')
            try:
                if not path.exists():self.storage.download_file(r2key,path)
            except Exception:path=None
        if path:
            try:
                with Image.open(path) as im:im.verify()
                if key=='sem':self.hashes[rid]=hashlib.sha256(path.read_bytes()).hexdigest()
                return path
            except Exception:pass
        self.missing.append({'point_id':rid,'asset':key});return None
    def clear_downloads(self):
        for p in self.folder.glob('*.jpg'):p.unlink(missing_ok=True)

def metadata(p,basis,criterion="co",n_threshold=2.40):
    c,n,o,src=feature_element_ratios(p.get('features') or {})
    return f"{condition(p)} | W{p.get('wafer')} {POSITIONS.get(integer(p.get('wafer')),'Other')} | P{p.get('point')} | {substrate(p)}",c,n,o,src

def point_page(doc,p,opts,assets):
    line,c,n,o,src=metadata(p,opts.cohort.basis,opts.cohort.criterion,opts.cohort.n_threshold)
    result=decision(p,opts.cohort.basis,opts.cohort.criterion,opts.cohort.n_threshold)
    sc=score(p,opts.cohort.criterion,opts.cohort.n_threshold)
    doc.page(f"{condition(p)}   W{p.get('wafer')} / P{p.get('point')}",f"{substrate(p)} | {POSITIONS.get(integer(p.get('wafer')),'Other')} | {result} | {opts.cohort.basis.upper()} decision basis")
    f=p.get('features') or {};polygons=(f.get('human_roi_polygons') or f.get('human_roi_polygon')) if opts.include_verification else None
    # Keep v23.9.14/15 one-page composition, adding N beside C/O without reducing the main SEM/EDS evidence.
    panels=[(.35,1.35,4.6,3.08,'sem','SEM original'+(' + saved Human ROI' if polygons else '')),(5.12,1.35,3.6,3.08,'eds_map','Full EDS original'),(8.9,1.35,4.05,5.32,'full_element_maps_original','Element maps original'),(.35,4.68,1.45,1.85,'c_map','C map'+(' + ROI' if polygons else '')),(1.92,4.68,1.45,1.85,'n_map','N map'+(' + ROI' if polygons else '')),(3.49,4.68,1.45,1.85,'o_map','O map'+(' + ROI' if polygons else ''))]
    for x,y,w,h,key,title in panels:
        doc.text(x,y-.2,w,.2,title,10,bold=True);doc.image(assets.get(p,key),x,y,w,h,polygons if key in ('sem','c_map','n_map','o_map') else None)
    doc.rect(5.12,4.48,3.6,2.2,LIGHT)
    lines=[f"Result: {result}",f"C ROI/Global: {textnum(c)}x",f"N ROI/Global: {textnum(n)}x",f"O ROI/Global: {textnum(o)}x",f"Score: {textnum(sc)} / 100"]
    if opts.include_verification:lines += [f"Human: {label(p.get('human_result')) or 'Not verified'}",f"Ratio source: {src}",f"Reviewed: {p.get('human_verified_at') or p.get('human_updated_at') or '-'}"]
    doc.text(5.28,4.6,3.26,1.95,'\n'.join(lines),9.2)
    doc.text(.35,6.78,12.5,.22,f"Point ID: {p.get('id')} | Batch: {p.get('project_id') or 'Unknown'}",8,color=MUTED)

def gallery_page(doc,key,pts,opts,assets,start,total):
    sub,res,cond,b=key
    band_label={'0-60':'0 <= score < 60','60-70':'60 <= score < 70','70-80':'70 <= score < 80','80-90':'80 <= score < 90','90-100':'90 <= score <= 100','missing':'Score unavailable'}.get(b,b)
    doc.page(f"Original SEM | {res} | {band_label}",f"{sub} | {cond} | {start+1}-{start+len(pts)} of {total} | {opts.cohort.basis.upper()} decision basis | Original stored SEM, no ROI or enhancement")
    cols,rows={10:(5,2),12:(4,3),16:(4,4),20:(5,4)}[opts.per_page];cw=12.6/cols;ch=5.75/rows
    for i,p in enumerate(pts):
        x=.35+(i%cols)*cw;y=1.26+(i//cols)*ch
        doc.image(assets.get(p,'sem'),x+.045,y,cw-.14,ch-.56)
        doc.text(x+.045,y+ch-.52,cw-.13,.16,f"{p.get('power')} {p.get('time')} | W{p.get('wafer')} P{p.get('point')}",8,bold=True)
        doc.text(x+.045,y+ch-.34,cw-.13,.15,f"Score {textnum(score(p,opts.cohort.criterion,opts.cohort.n_threshold))} | {str(p.get('id',''))[:10]}",7,color=MUTED)

def summary_pages(doc,records,opts):
    data=summary(records,opts.cohort)
    doc.page(opts.title,'Condition / location evaluation | W1 Corner, W4 Edge, W5 Middle | W9 reference evaluated separately')
    doc.text(.4,1.45,12,1.0,'Plasma activation and UV tape residue screening',27,bold=True)
    doc.text(.4,2.5,11.9,1.1,'SEM morphology and the selected C/N/O enrichment rule support a residue screening decision.\nResidue denotes the configured screening class; it is not source-specific chemical identification.',17,color=MUTED)
    st=stats(records,opts.cohort.basis,opts.cohort.criterion,opts.cohort.n_threshold)
    doc.table(['Selected points','Residue','Non-residue','Ambiguous','Human verified'],[[str(st[k]) for k in ['n','residue','non','ambiguous','verified']]],.4,3.9,[2.5]*5,row_h=.65,size=15)
    doc.text(.4,5.5,12,1.25,f"Scope: {opts.cohort.scope} | Decision basis: {opts.cohort.basis}\nRate = residue points / observed points. Ambiguous remains in denominator.\nPooled location rates reflect condition mix. Balanced rates equally weight common conditions.\nSelected-rule score is a screening index, not probability, chemical concentration or removal efficiency.",12,color=MUTED)
    for s in data['strata']:
        for start in range(0,len(s['conditions']),5):
            rows=s['conditions'][start:start+5];ws=[c['wafer'] for c in rows[0]['cells']] if rows else []
            doc.page('Plasma condition comparison',f"Substrate: {s['substrate']} | {opts.cohort.scope} | {opts.cohort.basis.upper()} basis | Rate [Residue / n]")
            headers=['Condition']+[f"W{w} {POSITIONS.get(w,'Other')}" for w in ws]+['R / n','Amb.','Spread']
            widths=[2.3]+[7.0/max(1,len(ws))]*len(ws)+[1.15,.9,1.25]
            data_rows=[]
            for r in rows:data_rows.append([r['condition']]+[f"{percent(c['rate'])}\n{c['residue']} / {c['n']}" for c in r['cells']]+[f"{r['residue']}/{r['n']}",str(r['ambiguous']),textnum(r['spread'])+' pp' if r['spread'] is not None else '-'])
            doc.table(headers,data_rows,.35,1.38,widths,row_h=.72,size=12)
            doc.text(.4,6.3,12.3,.6,'Missing observations are shown as -. Spread requires all selected wafer positions.\nRepeated W/P labels are not assumed to be paired physical specimens.',10,color=MUTED)
        doc.page('Wafer position across plasma conditions',f"Substrate: {s['substrate']} | Common conditions: {len(s['common_keys'])} | All selected conditions pooled vs equal-condition weighting")
        doc.table(['Position','N','Residue / n','Pooled','Balanced','Ambiguous'],[[f"W{p['wafer']} {p['label']}",str(p['n']),f"{p['residue']}/{p['n']}",percent(p['rate']),percent(p['balanced']),str(p['ambiguous'])] for p in s['positions']],.35,1.35,[2.5,1,2,2,2,3.1],row_h=.65,size=13)
        y=1.35+.65*(len(s['positions'])+1)+.35
        if len(s['positions'])<=4:
            for i,p in enumerate(s['positions']):
                yy=y+i*.42;doc.text(.45,yy,2,.2,f"W{p['wafer']} {p['label']}",11);doc.rect(2.8,yy,7,.17,LIGHT)
                if p['balanced'] is not None:doc.rect(2.8,yy,7*p['balanced']/100,.17,BLUE)
                doc.text(10.05,yy,2,.2,percent(p['balanced']),11)
        doc.text(.4,6.45,12.3,.4,'Balanced = mean of each common condition rate at this position. This is descriptive standardization, not a causal effect.',10,color=MUTED)

def comparison_page(doc,records,opts,assets):
    doc.page('Selected Point comparison','Original stored images | Columns are independent observations; scale bars retain original units')
    width=12.6/len(records);maps=list(dict.fromkeys(opts.comparison_maps));height=4.4/len(maps)
    for i,p in enumerate(records):
        x=.35+i*width;ident=p.get('research_identity') or {};_,c,n,o,_=metadata(p,opts.cohort.basis,opts.cohort.criterion,opts.cohort.n_threshold)
        doc.text(x+.06,1.22,width-.12,.22,f"{p.get('power')} {p.get('time')} | W{p.get('wafer')} P{p.get('point')}",11,bold=True)
        doc.text(x+.06,1.52,width-.12,.28,f"{ident.get('sample_id') or '-'} | Repeat {ident.get('repeat_index') or '-'} | {str(p.get('project_id',''))[:12]}",8,color=MUTED)
        for j,key in enumerate(maps):
            y=1.95+j*height;doc.text(x+.06,y-.15,width-.12,.15,key.replace('_map','').upper(),8,bold=True)
            doc.image(assets.get(p,key),x+.06,y,width-.12,height-.25)
        doc.text(x+.06,6.45,width-.12,.22,f"{decision(p,opts.cohort.basis,opts.cohort.criterion,opts.cohort.n_threshold)} | Score {textnum(score(p,opts.cohort.criterion,opts.cohort.n_threshold))}",9,bold=True)
        doc.text(x+.06,6.75,width-.12,.2,f"{substrate(p)} | C {textnum(c)}x | N {textnum(n)}x | O {textnum(o)}x",9,color=MUTED)

def _render_parts(records,opts,folder,storage=None,progress=lambda **x:None):
    folder=Path(folder);folder.mkdir(parents=True,exist_ok=True);assets=Assets(folder/'assets',storage)
    total=page_count(records,opts);outputs=[];manifest=[];done=0;max_pages=25 if opts.kind=='points' else 10
    doc=None;part=0
    def page_ready():
        nonlocal doc,part
        if doc is None:
            part+=1;doc=Document(opts.format,folder/f'{opts.kind}_{part:03d}.{opts.format}',footer_text=f"SEM / EDS screening | {criterion_rule_text(opts.cohort.criterion,opts.cohort.n_threshold)}");doc.page_offset=done
    def flush():
        nonlocal doc
        if doc is not None:outputs.append(doc.save());doc=None;assets.clear_downloads()
    if opts.kind=='summary':
        page_ready();summary_pages(doc,records,opts);done=total;flush()
    elif opts.kind=='comparison':
        page_ready();comparison_page(doc,records,opts,assets);done=1
        manifest.extend({'point_id':p['id'],'file':doc.out.name,'page':1,'column':i+1} for i,p in enumerate(records));flush()
    elif opts.kind=='points':
        for p in records:
            page_ready();point_page(doc,p,opts,assets);done+=1;progress(progress=round(done/total*95),message=f'Point {done}/{total}')
            manifest.append({'point_id':p['id'],'file':doc.out.name,'page':doc.number})
            if doc.number>=max_pages:flush()
        flush()
    else:
        for key,pts in gallery_groups(records,opts):
            for start in range(0,len(pts),opts.per_page):
                page_ready();a=pts[start:start+opts.per_page];gallery_page(doc,key,a,opts,assets,start,len(pts));done+=1
                for p in a:manifest.append({'point_id':p['id'],'file':doc.out.name,'page':doc.number,'group':list(key)})
                progress(progress=round(done/total*95),message=f'SEM gallery {done}/{total}')
                if doc.number>=max_pages:flush()
        flush()
    # Parts are temporary working files; the delivered document is always merged.
    from .report_merge import merge_documents
    final=folder/f'{opts.kind}.{opts.format}'
    progress(progress=96,message='최종 파일 통합 중')
    merge_documents(outputs,final,opts.format)
    offsets={};offset=0
    for part_file in outputs:
        offsets[part_file.name]=offset
        offset+=document_pages(part_file,opts.format)
    for entry in manifest:
        entry['page']+=offsets[entry['file']];entry['file']=final.name
    for part_file in outputs:part_file.unlink()
    outputs=[final]
    # Always deliver an audit manifest and original-SEM hashes alongside the outputs.
    audit={'created_at':datetime.now(timezone.utc).isoformat(),'options':opts.model_dump(),'point_ids':[p['id'] for p in records],'identities':{p['id']:p.get('research_identity') or {} for p in records},'n':len(records),'pages':total,'rule':criterion_rule_text(opts.cohort.criterion,opts.cohort.n_threshold),'source_note':'Original means stored SEM extracted from source PDF. No overlay, enhancement or crop added in gallery.','missing_assets':assets.missing,'sem_sha256':assets.hashes,'page_index':manifest}
    (folder/'manifest.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf8')
    with (folder/'points.csv').open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.writer(f);writer.writerow(['Point ID','Batch','Condition','Wafer','Position','Point','Substrate','Result','Human','Criterion','C ratio','N ratio','O ratio','Score','Score band','Sample ID','Repeat','Measured on','Batch label'])
        for p in records:
            _,c,n,o,_=metadata(p,opts.cohort.basis,opts.cohort.criterion,opts.cohort.n_threshold)
            row=[p['id'],p.get('project_id'),condition(p),p.get('wafer'),POSITIONS.get(integer(p.get('wafer')),'Other'),p.get('point'),substrate(p),decision(p,opts.cohort.basis,opts.cohort.criterion,opts.cohort.n_threshold),label(p.get('human_result')) or '',opts.cohort.criterion,c,n,o,score(p,opts.cohort.criterion,opts.cohort.n_threshold),band(p,opts.cohort.criterion,opts.cohort.n_threshold),*((p.get('research_identity') or {}).get(k,'') for k in ['sample_id','repeat_index','measured_on','batch_label'])]
            writer.writerow(["'"+v if isinstance(v,str) and v.startswith(('=','+','-','@','\t','\r')) else v for v in row])
    archive=folder/'report_bundle.zip'
    with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as z:
        for p in outputs+[folder/'manifest.json',folder/'points.csv']:z.write(p,p.name)
    return {'file':archive.name,'n':len(records),'pages':total,'parts':len(outputs),'documents':[p.name for p in outputs],'missing_assets':len(assets.missing)}


def document_pages(path,fmt):
    if fmt=='pdf':
        import fitz
        with fitz.open(path) as d:return len(d)
    from lxml import etree
    with zipfile.ZipFile(path) as z:
        root=etree.fromstring(z.read('ppt/presentation.xml'))
        return len(root.findall('{http://schemas.openxmlformats.org/presentationml/2006/main}sldIdLst/*'))

def render_export(records,opts,folder,storage=None,progress=lambda **x:None):
    if not opts.split_by_result or opts.kind in ('summary','comparison'):
        return _render_parts(records,opts,folder,storage,progress)
    import shutil
    folder=Path(folder);folder.mkdir(parents=True,exist_ok=True)
    groups=[(label,[p for p in records if decision(p,opts.cohort.basis,opts.cohort.criterion,opts.cohort.n_threshold)==label]) for label in RESULTS]
    groups=[(label,pts) for label,pts in groups if pts]
    results=[];audits=[];outputs=[]
    for index,(label,pts) in enumerate(groups):
        sub=folder/label
        result=_render_parts(pts,opts.model_copy(update={'split_by_result':False}),sub,storage,
            lambda **v:progress(progress=round((index+v.get('progress',0)/100)/len(groups)*99),message=f"{label} · {v.get('message','')}"))
        name=f'{opts.kind}_{label}.{opts.format}'
        shutil.move(str(sub/f'{opts.kind}.{opts.format}'),str(folder/name));outputs.append(folder/name)
        audit=json.loads((sub/'manifest.json').read_text(encoding='utf8'))
        for entry in audit['page_index']:entry['file']=name
        audits.append(audit);results.append(result)
        shutil.copyfile(sub/'points.csv',folder/f'points_{label}.csv');outputs.append(folder/f'points_{label}.csv')
        shutil.rmtree(sub)
    audit={**audits[0],'options':opts.model_dump(),'n':len(records),'pages':sum(x['pages'] for x in results),
        'point_ids':[p['id'] for p in records],
        'identities':{k:v for a in audits for k,v in a['identities'].items()},
        'sem_sha256':{k:v for a in audits for k,v in a['sem_sha256'].items()},
        'missing_assets':[e for a in audits for e in a['missing_assets']],
        'page_index':[e for a in audits for e in a['page_index']]}
    (folder/'manifest.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf8');outputs.append(folder/'manifest.json')
    with zipfile.ZipFile(folder/'report_bundle.zip','w',zipfile.ZIP_DEFLATED) as z:
        for p in outputs:z.write(p,p.name)
    return {'file':'report_bundle.zip','n':len(records),'pages':audit['pages'],'parts':len(groups),'documents':[p.name for p in outputs if p.suffix=='.'+opts.format],'missing_assets':len(audit['missing_assets'])}
