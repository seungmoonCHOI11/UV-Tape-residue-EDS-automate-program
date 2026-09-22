from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from reportlab.lib.pagesizes import landscape, A4
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader

def _add_box(slide, x,y,w,h,title):
    shape=slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.line.color.rgb=RGBColor(210,220,230)
    shape.fill.solid(); shape.fill.fore_color.rgb=RGBColor(255,255,255)
    tb=slide.shapes.add_textbox(Inches(x+.08), Inches(y+.05), Inches(w-.16), Inches(.2))
    tb.text_frame.paragraphs[0].text=title
    tb.text_frame.paragraphs[0].font.size=Pt(8); tb.text_frame.paragraphs[0].font.bold=True
    return shape

def export_ppt(records, out):
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    blank=prs.slide_layouts[6]
    for r in records:
        s=prs.slides.add_slide(blank)
        p=r
        s.shapes.add_textbox(Inches(.25),Inches(.12),Inches(9),Inches(.3)).text_frame.paragraphs[0].text=f"{p['power']}_{p['time']}_W{p['wafer']}_P{p['point']}"
        for shp in s.shapes:
            if hasattr(shp,"text_frame") and shp.text:
                shp.text_frame.paragraphs[0].font.size=Pt(18); shp.text_frame.paragraphs[0].font.bold=True
                shp.text_frame.paragraphs[0].font.color.rgb=RGBColor(14,41,69)
        subtitle=s.shapes.add_textbox(Inches(.25),Inches(.46),Inches(12),Inches(.2))
        subtitle.text_frame.paragraphs[0].text=f"{p['power']} / {p['time']} · W{p['wafer']} / P{p['point']} · {p['zone']} · {p.get('human_result') or p.get('ai_result') or 'Review'}"
        subtitle.text_frame.paragraphs[0].font.size=Pt(8)
        for x,y,w,h,key,title in [(.25,.88,6,2.3,"sem","SEM"),(6.45,.88,6,2.3,"spectrum","EDS Spectrum"),(.25,3.36,6,2.3,"eds_map","EDS Map"),(6.45,3.36,6,2.3,"element_maps","Element Maps")]:
            _add_box(s,x,y,w,h,title)
            img=p.get("assets",{}).get(key)
            if img and Path(img).exists():
                s.shapes.add_picture(img, Inches(x+.15), Inches(y+.3), width=Inches(w-.3), height=Inches(h-.45))
        _add_box(s,.25,5.84,7.7,1.25,"EDS / Element Data")
        t=s.shapes.add_textbox(Inches(.4),Inches(6.18),Inches(7.3),Inches(.65))
        f=p.get("features",{})
        t.text_frame.paragraphs[0].text=f"Si —   C enrichment {f.get('c_enrichment',0)}%   N —   O enrichment {f.get('o_enrichment',0)}%\nC coverage {f.get('c_coverage',0)}%   O coverage {f.get('o_coverage',0)}%   Wafer W{p['wafer']}   Point P{p['point']}"
        t.text_frame.paragraphs[0].font.size=Pt(9)
        _add_box(s,8.15,5.84,4.3,1.25,"Result")
        t=s.shapes.add_textbox(Inches(8.35),Inches(6.3),Inches(3.9),Inches(.45))
        t.text_frame.paragraphs[0].text=p.get("human_result") or p.get("ai_result") or "Review"
        t.text_frame.paragraphs[0].font.size=Pt(21); t.text_frame.paragraphs[0].font.bold=True
        t.text_frame.paragraphs[0].alignment=2
    prs.save(out)
    return out

def export_pdf(records,out):
    c=canvas.Canvas(str(out),pagesize=landscape(A4)); W,H=landscape(A4)
    for idx,p in enumerate(records):
        c.setFont("Helvetica-Bold",15); c.setFillColorRGB(.05,.16,.27)
        c.drawString(22,H-25,f"{p['power']}_{p['time']}_W{p['wafer']}_P{p['point']}")
        c.setFont("Helvetica",7); c.drawString(22,H-37,f"{p['power']} / {p['time']} · W{p['wafer']} / P{p['point']} · {p['zone']}")
        c.setStrokeColorRGB(.2,.36,.51); c.line(22,H-45,W-22,H-45)
        boxes=[(22,H-215,330,155,"SEM","sem"),(365,H-215,330,155,"EDS Spectrum","spectrum"),(22,H-380,330,155,"EDS Map","eds_map"),(365,H-380,330,155,"Element Maps","element_maps")]
        for x,y,w,h,title,key in boxes:
            c.setStrokeColorRGB(.82,.86,.89); c.rect(x,y,w,h)
            c.setFillColorRGB(.05,.18,.28); c.setFont("Helvetica-Bold",7); c.drawString(x+7,y+h-12,title)
            img=p.get("assets",{}).get(key)
            if img and Path(img).exists():
                c.drawImage(ImageReader(img),x+7,y+7,w-14,h-24,preserveAspectRatio=True,anchor='c')
        c.rect(22,52,420,70); c.rect(452,52,283,70)
        c.setFillColorRGB(.05,.1,.15); c.setFont("Helvetica-Bold",7); c.drawString(28,108,"EDS / Element Data")
        f=p.get("features",{})
        c.setFont("Helvetica",7); c.drawString(28,92,f"C enrichment: {f.get('c_enrichment',0)}%   O enrichment: {f.get('o_enrichment',0)}%")
        c.drawString(28,78,f"C coverage: {f.get('c_coverage',0)}%   O coverage: {f.get('o_coverage',0)}%   W{p['wafer']} / P{p['point']}")
        c.setFont("Helvetica-Bold",7); c.drawString(460,108,"Result")
        c.setFont("Helvetica-Bold",18); c.drawCentredString(593,76,p.get("human_result") or p.get("ai_result") or "Review")
        c.setFont("Helvetica",6); c.drawRightString(W-25,20,f"{idx+1} / {len(records)}")
        c.showPage()
    c.save(); return out
