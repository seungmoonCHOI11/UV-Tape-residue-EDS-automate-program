import hashlib,json,math,time,zipfile,threading
from pathlib import Path
import pytest
from PIL import Image,ImageDraw
from pptx import Presentation
import fitz
from app.research import Cohort,ExportOptions,select,summary,band,score,gallery_groups,page_count
from app.research_reports import render_export,fit

def point(i,w=1,power='150W',result='Residue',sub='SiCN',batch='b1',score_value=None):
    f={'result':result}
    if score_value is not None:f['residue_score']=score_value/100
    return {'id':f'p{i}','project_id':batch,'power':power,'time':'30s','wafer':w,'point':i%9+1,'substrate_type':sub,'features':f,'assets':{}}

def images(tmp_path):
    im=Image.new('RGB',(900,600),(40,55,70));d=ImageDraw.Draw(im);d.rectangle((1,1,898,598),outline=(250,200,50),width=6);d.text((35,35),'SYNTHETIC QA IMAGE - NOT EXPERIMENTAL DATA',fill='white');d.rectangle((640,530,830,542),fill='white');p=tmp_path/'raw.jpg';im.save(p,quality=90);return p

def test_scope_and_verified_and_score_edges():
    a=[point(i,w=w,sub=sub,score_value=sc) for i,(w,sub,sc) in enumerate([(1,'SiCN',70),(4,'SiCN',79.9),(5,'SiCN',80),(9,'Si',90),(5,'SiCN',100)])]
    assert [band(p) for p in a]==['70-80','70-80','80-90','90-100','90-100']
    assert len(select(a,Cohort()))==4
    assert len(select(a,Cohort(scope='reference',substrates=['Si'])))==1
    a[0]['human_result']='Non-residue'
    assert len(select(a,Cohort(results=['Non-residue'],review='verified')))==1
    assert len(select(a,Cohort(results=['Residue'],basis='auto')))==4
    assert select(a,Cohort(conditions=[]))==[]
    assert band(point(55))=='missing'

def test_pooled_vs_balanced_and_missing_are_not_zero():
    # Different condition mix reverses pooled location comparison; balanced recovers equal weights.
    a=[];i=0
    for power,w,n,r in [('150W',1,10,9),('350W',1,100,10),('150W',4,100,80),('350W',4,10,0)]:
        for j in range(n):i+=1;a.append(point(i,w,power,'Residue' if j<r else 'Non-residue'))
    s=summary(a,Cohort(wafers=[1,4]))['strata'][0]
    p1,p4=s['positions'];assert p1['rate']<p4['rate'];assert p1['balanced']>p4['balanced'];assert p1['common_conditions']==2
    s=summary(a,Cohort())['strata'][0];assert s['positions'][2]['rate'] is None;assert all(p['balanced'] is None for p in s['positions']);assert all(c['spread'] is None for c in s['conditions'])

def test_substrates_and_duplicate_coordinates_are_separate():
    a=[point(1),point(1),point(10),point(20,sub='Si')]
    selected=select(a,Cohort());assert len(selected)==3
    strata=summary(selected,Cohort())['strata'];assert len(strata)==2
    si=next(s for s in strata if s['substrate']=='SiCN');assert si['conditions'][0]['slots']==1;assert si['conditions'][0]['duplicates']==1

def test_gallery_does_not_mix_classes_bands_or_substrates():
    a=[point(1,score_value=70),point(2,score_value=80),point(3,result='Non-residue',score_value=80),point(4,sub='Si',score_value=80)]
    opts=ExportOptions(kind='gallery',group_conditions=False)
    assert len(gallery_groups(a,opts))==4
    assert page_count(a,opts)==4

@pytest.mark.parametrize('fmt',['pptx','pdf'])
@pytest.mark.parametrize('per_page',[10,12,16,20])
def test_gallery_pages_raw_bytes_and_aspect(tmp_path,fmt,per_page):
    raw=images(tmp_path);digest=hashlib.sha256(raw.read_bytes()).hexdigest()
    a=[point(i,score_value=75) for i in range(per_page+1)]
    for p in a:p['assets']={'sem':str(raw),'sem_residue_overlay':'MUST_NOT_USE'}
    opts=ExportOptions(kind='gallery',format=fmt,per_page=per_page)
    folder=tmp_path/'out';result=render_export(a,opts,folder)
    assert result['pages']==2
    manifest=json.loads((folder/'manifest.json').read_text());assert len(manifest['page_index'])==len(a);assert set(manifest['sem_sha256'].values())=={digest}
    assert hashlib.sha256(raw.read_bytes()).hexdigest()==digest
    file=next(folder.glob(f'gallery_*.{fmt}'))
    if fmt=='pptx':
        prs=Presentation(file);assert len(prs.slides)==2
        pics=[sh for sl in prs.slides for sh in sl.shapes if sh.shape_type==13];assert len(pics)==len(a)
        assert all(abs(sh.width/sh.height-1.5)<.001 for sh in pics);assert all(sh.crop_top==sh.crop_bottom==sh.crop_left==sh.crop_right==0 for sh in pics)
        assert all(hashlib.sha256(sh.image.blob).hexdigest()==digest for sh in pics)
    else:
        doc=fitz.open(file);assert len(doc)==2;assert 'Original SEM' in doc[0].get_text()

def test_missing_raw_sem_never_uses_overlay(tmp_path):
    p=point(1,score_value=75);p['assets']={'sem_residue_overlay':str(images(tmp_path))}
    render_export([p],ExportOptions(kind='gallery',format='pdf'),tmp_path/'out')
    m=json.loads((tmp_path/'out/manifest.json').read_text());assert m['missing_assets']==[{'point_id':'p1','asset':'sem'}];assert m['sem_sha256']=={}

def test_points_verification_and_split(tmp_path):
    raw=images(tmp_path);a=[point(i,score_value=75) for i in range(26)]
    for p in a:p['assets']={k:str(raw) for k in ['sem','eds_map','full_element_maps_original','c_map','o_map']};p['human_result']='Residue';p['features']['human_roi_polygon']=[{'x':.2,'y':.2},{'x':.4,'y':.2},{'x':.3,'y':.4}]
    r=render_export(a,ExportOptions(kind='points',include_verification=True),tmp_path/'out');assert r['parts']==2;assert r['pages']==26
    prs=Presentation(tmp_path/'out/points_001.pptx');text=' '.join(sh.text for sh in prs.slides[0].shapes if sh.has_text_frame);assert 'Human: Residue' in text;assert 'saved Human ROI' in text
    m=json.loads((tmp_path/'out/manifest.json').read_text());assert len(m['page_index'])==26

def test_summary_pagination(tmp_path):
    a=[point(i,w,power=f'{150+20*i}W') for i in range(12) for w in [1,4,5]]
    # Unique IDs matter; give each sample its own ID.
    for i,p in enumerate(a):p['id']=str(i)
    opts=ExportOptions(kind='summary',format='pdf');r=render_export(a,opts,tmp_path/'out');doc=fitz.open(tmp_path/'out/summary_001.pdf');assert len(doc)==5;assert r['pages']==5

def test_paged_database_keeps_more_than_response_cap():
    from app.db import _paged
    from types import SimpleNamespace
    class Query:
        def range(self,a,b):self.a=a;self.b=b;return self
        def execute(self):return SimpleNamespace(data=list(range(1257))[self.a:self.b+1])
    assert len(_paged(Query()))==1257

def test_report_api_queues_behind_analysis(tmp_path):
    from app.report_api import install_report_routes
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    app=FastAPI();slot=threading.Lock();slot.acquire();raw=images(tmp_path);p=point(1,score_value=75);p['assets']={'sem':str(raw)}
    install_report_routes(app,lambda:[p],slot,tmp_path);client=TestClient(app)
    r=client.post('/api/research/reports',json={'kind':'gallery','format':'pdf'});assert r.status_code==200;job=r.json()['id']
    assert client.get(f'/api/research/reports/{job}').json()['status']=='queued'
    assert client.get(f'/api/research/reports/{job}/download').status_code==409
    slot.release()
    for _ in range(150):
        state=client.get(f'/api/research/reports/{job}').json()
        if state['status'] in ['completed','failed']:break
        time.sleep(.1)
    assert state['status']=='completed',state
    assert client.get(f'/api/research/reports/{job}/download').content[:2]==b'PK'
    assert client.post('/api/research/reports',json={'cohort':{'conditions':[]}}).status_code==422
    assert client.post('/api/research/reports',json={'per_page':21}).status_code==422
