import json,zipfile,posixpath
import numpy as np,pytest
from PIL import Image
from lxml import etree as ET
from app.point_worker import analytical_mask,exclude_element_annotations,global_roi_metrics
from app.research import ExportOptions
from app.research_reports import render_export,document_pages

def test_annotation_pixels_never_change_metrics_and_display_preserves_field():
    im=np.random.default_rng(43).integers(5,90,(160,200,3),dtype=np.uint8);roi=np.zeros((160,200),np.uint8);roi[30:70,50:100]=255
    am=analytical_mask(im.shape)>0
    for channel in ['C','O']:
        original=global_roi_metrics(im,roi,channel)
        altered=im.copy();altered[~am]=255
        assert global_roi_metrics(altered,roi,channel)==original
    shown=exclude_element_annotations(im)
    assert np.array_equal(shown[am],im[am])
    assert np.array_equal(shown[~am,0],shown[~am,1])
    assert np.array_equal(shown[~am,1],shown[~am,2])

def records(tmp,n=162):
    image=tmp/'source.png';Image.new('RGB',(120,90),(80,90,100)).save(image)
    return [dict(id=f'p{i:03}',project_id='repeat1',power='150W',time='30s',wafer=[1,4,5][i%3],point=i%9+1,
        human_result=['Residue','Non-residue','Ambiguous'][i%3],substrate_type='SiCN',
        features={'c_roi_global_ratio':3,'o_roi_global_ratio':4},
        assets={k:str(image) for k in ['sem','eds_map','full_element_maps_original','c_map','o_map']}) for i in range(n)]

def validate_ppt(path,count):
    from pptx import Presentation
    d=Presentation(path);assert len(d.slides)==count
    for slide in d.slides:
        assert len(slide.shapes)>5
        for pic in [s for s in slide.shapes if s.shape_type==13]:assert pic.image.blob
    with zipfile.ZipFile(path) as z:
        names=set(z.namelist())
        for name in names:
            if not name.endswith('.rels'):continue
            basedir=posixpath.dirname(posixpath.dirname(name))
            for rel in ET.fromstring(z.read(name)):
                if rel.get('TargetMode')=='External':continue
                target=rel.get('Target')
                resolved=target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join(basedir,target))
                assert resolved in names,(name,resolved)

@pytest.mark.parametrize('fmt',['pdf','pptx'])
def test_162_point_report_single_document_and_manifest(tmp_path,fmt):
    pts=records(tmp_path);out=tmp_path/'out'
    result=render_export(pts,ExportOptions(kind='points',format=fmt),out)
    assert result['parts']==1 and result['pages']==162
    path=out/f'points.{fmt}'
    assert document_pages(path,fmt)==162
    audit=json.loads((out/'manifest.json').read_text())
    assert [e['point_id'] for e in audit['page_index']]==[p['id'] for p in pts]
    assert [e['page'] for e in audit['page_index']]==list(range(1,163))
    assert {e['file'] for e in audit['page_index']}=={path.name}
    with zipfile.ZipFile(out/'report_bundle.zip') as z:assert len([x for x in z.namelist() if x.endswith('.'+fmt)])==1
    if fmt=='pptx':validate_ppt(path,162)
    else:
        import fitz
        with fitz.open(path) as d:
            for i in [0,24,25,149,161]:assert f'p{i:03}' in d[i].get_text()

@pytest.mark.parametrize('fmt',['pdf','pptx'])
def test_result_split_and_empty_class_omission(tmp_path,fmt):
    pts=[p for p in records(tmp_path,52) if p['human_result']!='Ambiguous']
    result=render_export(pts,ExportOptions(kind='points',format=fmt,split_by_result=True),tmp_path/'out')
    assert result['parts']==2 and len(result['documents'])==2
    assert all('Ambiguous' not in x for x in result['documents'])
    audit=json.loads((tmp_path/'out/manifest.json').read_text())
    assert len(audit['page_index'])==len(pts)
    for label in ['Residue','Non-residue']:
        entries=[e for e in audit['page_index'] if e['file']==f'points_{label}.{fmt}']
        assert [e['page'] for e in entries]==list(range(1,len(entries)+1))

@pytest.mark.parametrize('fmt',['pdf','pptx'])
def test_gallery_and_summary_single_document(tmp_path,fmt):
    pts=records(tmp_path)
    for kind in ['gallery','summary']:
        out=tmp_path/kind
        result=render_export(pts,ExportOptions(kind=kind,format=fmt,group_scores=False,group_conditions=False,per_page=10),out)
        assert result['parts']==1
        assert document_pages(out/f'{kind}.{fmt}',fmt)==result['pages']

def test_existing_map_display_does_not_analyze_or_modify_values(tmp_path,monkeypatch):
    import copy,cv2
    from fastapi.testclient import TestClient
    from app import main
    raw=np.zeros((150,180,3),np.uint8);raw[:,:,2]=170
    source=tmp_path/'c.jpg';cv2.imwrite(str(source),raw)
    point={'id':'display-test','features':{'c_roi_global_ratio':2.7,'o_roi_global_ratio':3.4},'human_result':'Residue','assets':{'c_map_enhanced_overlay':str(source)}}
    snapshot=copy.deepcopy(point)
    monkeypatch.setattr(main,'OUTPUT',tmp_path)
    monkeypatch.setattr(main,'RECORDS',{'display-test':point})
    monkeypatch.setattr(main.db,'configured',lambda:False)
    def forbidden(*a,**k):raise AssertionError('Display must not reanalyze')
    monkeypatch.setattr(main,'_dynamic_asset',forbidden)
    client=TestClient(main.app)
    for _ in range(2):
        response=client.get('/api/assets/display-test/c_map_enhanced_overlay')
        assert response.status_code==200
        im=cv2.imdecode(np.frombuffer(response.content,np.uint8),cv2.IMREAD_COLOR)
        assert (im[140,:,0]==im[140,:,2]).all()
        assert (im[50,10:100,2]>im[50,10:100,0]).all()
    assert point==snapshot

def test_direct_download_only_allows_completed_document(tmp_path):
    import threading
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from app.report_api import install_report_routes
    app=FastAPI();install_report_routes(app,lambda:[],threading.Lock(),tmp_path)
    folder=tmp_path/'_research_reports'/('a'*32);folder.mkdir()
    (folder/'points.pdf').write_bytes(b'%PDF-test')
    (folder/'result.json').write_text(json.dumps({'documents':['points.pdf']}))
    client=TestClient(app)
    response=client.get('/api/research/reports/'+('a'*32)+'/download?document=points.pdf')
    assert response.status_code==200 and response.content==b'%PDF-test'
    assert client.get('/api/research/reports/'+('a'*32)+'/download?document=worker.log').status_code==404
