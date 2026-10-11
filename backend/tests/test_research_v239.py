import json, threading, time
from types import SimpleNamespace
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from pptx import Presentation
import fitz
from app.research import Cohort, ExportOptions
from app.research_ai import install_ai_routes, interpret, InterpretationError, explain_error
from app.research_identity import install_identity_routes
from app.research_reports import render_export
from test_research_v238 import point, images


def test_identity_preserves_analysis_metadata_and_conflict():
    meta={'files':['original.pdf'],'job_state':{'status':'completed'},'repeat_no':7}
    projects={'b1':{'description':json.dumps(meta)}}
    slot=threading.Lock(); app=FastAPI()
    install_identity_routes(app,SimpleNamespace(configured=lambda:False),projects,slot,threading.Lock())
    client=TestClient(app)
    r=client.put('/api/projects/b1/research-identity',json={'sample_id':' S-001 ','repeat_index':2})
    assert r.status_code==200
    saved=json.loads(projects['b1']['description'])
    assert all(saved[k]==v for k,v in meta.items())
    assert saved['research_identity']['sample_id']=='S-001'
    assert client.put('/api/projects/b1/research-identity',json={'sample_id':'stale'}).status_code==409
    slot.acquire()
    assert client.put('/api/projects/b1/research-identity',json={}).status_code==409
    slot.release()
    assert client.put('/api/projects/b1/research-identity',json={'repeat_index':0}).status_code==422


@pytest.mark.parametrize('fmt',['pdf','pptx'])
@pytest.mark.parametrize('n',[2,4])
def test_comparison_single_page_raw_maps_and_order(tmp_path,fmt,n):
    raw=images(tmp_path); records=[point(i,w=9 if i==0 else 1) for i in reversed(range(n))]
    for p in records:
        p['assets']={k:str(raw) for k in ['sem','c_map','o_map']}
        p['research_identity']={'sample_id':'SPECIMEN-01','repeat_index':2}
    opts=ExportOptions(kind='comparison',format=fmt,cohort=Cohort(scope='all'),comparison_ids=[p['id'] for p in records],comparison_maps=['sem','c_map','o_map','n_map'])
    folder=tmp_path/'out'; r=render_export(records,opts,folder)
    assert r['pages']==1
    m=json.loads((folder/'manifest.json').read_text())
    assert [p['point_id'] for p in m['page_index']]==opts.comparison_ids
    assert [p['column'] for p in m['page_index']]==list(range(1,n+1))
    assert len(m['missing_assets'])==n
    assert all(p['asset']=='n_map' for p in m['missing_assets'])
    file=folder/f'comparison_001.{fmt}'
    if fmt=='pdf': assert len(fitz.open(file))==1
    else:
        prs=Presentation(file); assert len(prs.slides)==1
        pics=[sh for sh in prs.slides[0].shapes if sh.shape_type==13]
        assert len(pics)==n*3
        assert all(sh.image.blob==raw.read_bytes() and sh.crop_top==sh.crop_bottom==sh.crop_left==sh.crop_right==0 for sh in pics)


def test_comparison_api_rejects_missing_or_duplicate_ids(tmp_path):
    from app.report_api import install_report_routes
    app=FastAPI();slot=threading.Lock();slot.acquire()
    install_report_routes(app,lambda:[point(1),point(2)],slot,tmp_path)
    client=TestClient(app)
    for ids in [['p1'],['p1','p1'],['p1','absent']]:
        assert client.post('/api/research/reports',json={'kind':'comparison','comparison_ids':ids}).status_code==422
    slot.release()


def configure(monkeypatch):
    monkeypatch.setenv('OPENAI_API_KEY','synthetic-key-not-real')
    monkeypatch.setenv('OPENAI_ANALYSIS_MODEL','synthetic-model')


def test_ai_provider_contract_without_network(monkeypatch):
    configure(monkeypatch); calls={}
    payload={'summary':'집계 해석','key_findings':[],'condition_trends':[],'anomalies':[],'next_steps':[],'caveats':[]}
    def create(**kwargs):calls.update(kwargs);return SimpleNamespace(output_text=json.dumps(payload))
    def factory(**kwargs):
        assert kwargs=={'timeout':45.0,'max_retries':0}
        return SimpleNamespace(responses=SimpleNamespace(create=create))
    p=point(1);p['assets']={'sem':'PRIVATE_IMAGE_MUST_NOT_SEND'}
    r=interpret([p],Cohort(),factory)
    assert r['source']['n']==1 and r['source']['model']=='synthetic-model'
    assert 'PRIVATE_IMAGE' not in calls['input'] and calls['model']=='synthetic-model'
    payload['summary']=''
    with pytest.raises(InterpretationError,match='완성된 해석'):interpret([p],Cohort(),factory)


@pytest.mark.parametrize('status,code',[(401,'authentication'),(403,'permission'),(404,'model'),(429,'quota'),(500,'provider')])
def test_ai_provider_error_mapping(status,code):
    assert explain_error(SimpleNamespace(status_code=status)).code==code


def test_ai_jobs_scope_busy_and_safe_failure(monkeypatch):
    configure(monkeypatch); app=FastAPI(); gate=threading.Event();seen=[]
    def runner(records,cohort):
        seen.extend(records);gate.wait(2)
        raise InterpretationError('quota','사용 한도 확인')
    install_ai_routes(app,lambda:[point(1),point(2,w=9)],runner)
    client=TestClient(app)
    start=client.post('/api/research/ai',json={'scope':'reference'});assert start.status_code==200
    assert client.post('/api/research/ai',json={}).status_code==409
    gate.set()
    for _ in range(100):
        state=client.get('/api/research/ai/'+start.json()['id']).json()
        if state['status']!='running':break
        time.sleep(.01)
    assert state['status']=='failed' and state['code']=='quota'
    assert [p['wafer'] for p in seen]==[9]
    assert client.get('/api/research/ai/missing').status_code==404
    monkeypatch.delenv('OPENAI_API_KEY')
    assert client.post('/api/research/ai',json={}).status_code==503


def test_project_pagination_respects_limit(monkeypatch):
    from app import db
    class Query:
        def table(self,*args):return self
        def select(self,*args):return self
        def order(self,*args,**kwargs):return self
        def range(self,a,b):self.a,self.b=a,b;return self
        def execute(self):return SimpleNamespace(data=list(range(1257))[self.a:self.b+1])
    monkeypatch.setattr(db,'get_client',lambda:Query())
    assert len(db.get_projects(1200))==1200
    assert len(db.get_projects(10000))==1257
